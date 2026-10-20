#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把实验3.4 的真实截图填进下发模板，产出交付版报告。

模板中 19 个空白段落即插图槽位，逐个填图并在其后插入「图 3-54 …」式图题，
编号接着实验3.3 报告的 图 3-53 往下排。图题 / 学号行格式与已交付的
实验3.1、3.2、3.3 报告保持一致（宋体、居中）。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

WORK = Path(__file__).resolve().parent
SHOTS = WORK / "shots"
TEMPLATE = WORK / "template.docx"
OUTPUT = WORK / "out.docx"

STUDENT = "学号：__________    班级：________    姓名：________"
FIRST_FIG = 54          # 实验3.3 交付到 图 3-53，这里接着排

# 模板段落索引 -> (截图文件, 图题)
SLOTS = [
    (6,  "01-apt-netperf.png",      "在交换机 1 上执行 apt install netperf 安装 netperf"),
    (8,  "02-ifconfig-switch.png",  "以 root 用户登录交换机 1，执行 ifconfig 查看其 IP 地址"),
    (10, "03-ovs-delbr.png",        "ovs-vsctl show 查看原有网桥，ovs-vsctl del-br br0 删除并查看确认"),
    (12, "04-ifconfig-host.png",    "以 root 用户登录主机 1，执行 ifconfig 查看其 IP 地址"),
    (13, "05-ping-replies.png",     "主机 1 ping 交换机 1，测试两者之间的连通性"),
    (14, "06-ping-stats.png",       "ping 命令的统计结果，丢包率为 0"),
    (17, "07-ps-ovs.png",           "ps -ef | grep ovs 查看交换机 1 上的 OVS 进程"),
    (19, "08-addbr-addport.png",    "创建网桥 br0，并将 ens33 网卡挂载到 br0 上"),
    (21, "09-br0-ifconfig.png",     "将 ens33 的 IP 地址赋给 br0，ifconfig 查看配置结果"),
    (23, "10-netserver.png",        "将 OVS 交换机作为 netperf 服务器，启动 netserver 并查看进程"),
    (25, "11-apt-netperf-host.png", "在主机 1 上执行 apt install netperf 安装 netperf"),
    (27, "12-netperf-before.png",   "未设置 QoS 时 netperf 测得的 UDP 吞吐量"),
    (30, "13-qos-set.png",          "在交换机 1 上设置 ens33 的 ingress_policing_rate 和 ingress_policing_burst"),
    (32, "14-netperf-after1.png",   "设置 QoS 后第一次 netperf 测试结果"),
    (33, "15-netperf-after2.png",   "设置 QoS 后第二次 netperf 测试结果"),
    (37, "16-ovs-show-scene2.png",  "场景二：ovs-vsctl show 查看虚拟交换机信息"),
    (40, "17-list-port-br0.png",    "ovs-vsctl list port br0 查看 br0 端口的 UUID"),
    (41, "18-list-port-ens33.png",  "ovs-vsctl list port ens33 查看 ens33 端口的 UUID"),
    (43, "19-mirror.png",           "创建 mirror，将发往 ens33 和从 ens33 发出的数据包全部定向到 br0 端口"),
]

NOTE_AFTER_FIG = 68
NOTE_TEXT = ("结果分析：未设置 QoS 时，主机 1 与交换机 1 之间的 UDP 吞吐量约为 682/673 Mbit/s。"
             "在 ens33 上设置 ingress_policing_rate=100000（约 100 Mbit/s）、"
             "ingress_policing_burst=50000 后，客户端仍按原速率发送，但服务器端实际接收速率"
             "降至 5.55 Mbit/s 和 0 Mbit/s，说明超出承诺速率的报文已被 OVS 入口 policer 丢弃，"
             "QoS 限速功能生效。")


def _eastasia(run, font: str):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def caption_para(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    _eastasia(run, "宋体")
    return p


def note_para(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    _eastasia(run, "宋体")
    return p


def main() -> None:
    doc = Document(str(TEMPLATE))

    sec = doc.sections[0]
    usable = sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm
    print(f"正文宽度 {usable:.2f} cm")

    # --- 标题下补 学号/班级/姓名（模板第 1 段是空段）---------------------
    head = doc.paragraphs[1]
    for r in list(head.runs):
        r._element.getparent().remove(r._element)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(STUDENT)
    run.font.size = Pt(12)
    run.font.bold = True
    _eastasia(run, "宋体")

    # --- 逐个槽位填图 + 图题 -------------------------------------------
    paras = doc.paragraphs
    figure = FIRST_FIG
    for idx, name, cap in SLOTS:
        img = SHOTS / name
        if not img.exists():
            raise SystemExit(f"[!] 缺图：{img}")

        slot = paras[idx]
        if slot.text.strip():
            raise SystemExit(f"[!] 第 {idx} 段不是空槽位：{slot.text[:30]!r}")
        slot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        slot.paragraph_format.space_before = Pt(8)
        slot.paragraph_format.space_after = Pt(0)
        pic = slot.add_run().add_picture(str(img), width=Cm(usable))
        print(f"  {idx:>3} <- {name}  ({pic.width.cm:.2f}x{pic.height.cm:.2f} cm)")

        cap_par = caption_para(doc, f"图 3-{figure}  {cap}")
        slot._element.addnext(cap_par._element)

        if figure == NOTE_AFTER_FIG:
            note = note_para(doc, NOTE_TEXT)
            cap_par._element.addnext(note._element)

        figure += 1

    doc.save(str(OUTPUT))
    print(f"[OK] {OUTPUT}")
    print(f"     共插入 {figure - FIRST_FIG} 张图，编号 图 3-{FIRST_FIG} ~ 图 3-{figure - 1}")


if __name__ == "__main__":
    main()
